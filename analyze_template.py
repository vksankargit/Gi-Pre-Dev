import openpyxl

wb = openpyxl.load_workbook('01 Annual Plan Template.xlsx')

for sheet_name in wb.sheetnames:
    print(f'\n{"="*60}')
    print(f'Sheet: {sheet_name}')
    print("="*60)
    ws = wb[sheet_name]

    # Print first 5 rows
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i <= 5:
            if any(cell is not None for cell in row):
                print(f"Row {i}: {row}")
        else:
            break
